#!/usr/bin/env python3
"""Build a clean Word paper skeleton while keeping the supplied manual as reference."""

from __future__ import annotations

import argparse
import shutil
import tempfile
import uuid
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


def set_east_asia_font(style, name: str) -> None:
    style.font.name = "Times New Roman"
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)


def set_run_fonts(run, east_asia: str = "Source Han Serif CN") -> None:
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), east_asia)


def set_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instruction, separate, text, end):
        run._r.append(element)


def clear_body(document: Document) -> None:
    """Remove the supplied manual's content while retaining its document package."""
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def paragraph_style(document: Document, name: str):
    try:
        return document.styles[name]
    except KeyError:
        return document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)


def source_han_regular(font_zip: Path) -> bytes:
    if not font_zip.is_file():
        raise FileNotFoundError(font_zip)
    with ZipFile(font_zip) as archive:
        candidates = [
            name
            for name in archive.namelist()
            if name.endswith("SourceHanSerifCN-Regular.otf")
        ]
        if len(candidates) != 1:
            raise RuntimeError("SourceHanSerifCN-Regular.otf not found uniquely in LaTeX zip")
        return archive.read(candidates[0])


def obfuscate_odttf(font_data: bytes, font_key: uuid.UUID) -> bytes:
    """Apply the ECMA-376 embedded-font XOR transform to the first 32 bytes."""
    data = bytearray(font_data)
    key = bytes.fromhex(font_key.hex)[::-1]
    for index in range(min(32, len(data))):
        data[index] ^= key[index % 16]
    return bytes(data)


def embed_source_han(output: Path, font_zip: Path) -> None:
    """Embed the full OFL-licensed Source Han Serif font into a DOCX package."""
    ct_ns = "http://schemas.openxmlformats.org/package/2006/content-types"
    rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    w_ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    r_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    font_key = uuid.uuid4()
    font_payload = obfuscate_odttf(source_han_regular(font_zip), font_key)
    font_member = "word/fonts/SourceHanSerifCN-Regular.odttf"
    rels_member = "word/_rels/fontTable.xml.rels"
    relationship_id = "rIdSourceHanSerifCNRegular"

    with ZipFile(output, "r") as source:
        members = {info.filename: source.read(info.filename) for info in source.infolist()}

    content_types = ET.fromstring(members["[Content_Types].xml"])
    default_tag = f"{{{ct_ns}}}Default"
    if not any(node.get("Extension") == "odttf" for node in content_types.findall(default_tag)):
        ET.SubElement(
            content_types,
            default_tag,
            {
                "Extension": "odttf",
                "ContentType": "application/vnd.openxmlformats-officedocument.obfuscatedFont",
            },
        )
    ET.register_namespace("", ct_ns)
    members["[Content_Types].xml"] = ET.tostring(
        content_types, encoding="utf-8", xml_declaration=True
    )

    font_table = ET.fromstring(members["word/fontTable.xml"])
    font_tag = f"{{{w_ns}}}font"
    for node in list(font_table.findall(font_tag)):
        if node.get(f"{{{w_ns}}}name") == "Source Han Serif CN":
            font_table.remove(node)
    font = ET.SubElement(font_table, font_tag, {f"{{{w_ns}}}name": "Source Han Serif CN"})
    ET.SubElement(font, f"{{{w_ns}}}family", {f"{{{w_ns}}}val": "roman"})
    ET.SubElement(font, f"{{{w_ns}}}pitch", {f"{{{w_ns}}}val": "variable"})
    ET.SubElement(
        font,
        f"{{{w_ns}}}embedRegular",
        {
            f"{{{r_ns}}}id": relationship_id,
            f"{{{w_ns}}}fontKey": "{" + str(font_key).upper() + "}",
        },
    )
    ET.register_namespace("w", w_ns)
    ET.register_namespace("r", r_ns)
    members["word/fontTable.xml"] = ET.tostring(
        font_table, encoding="utf-8", xml_declaration=True
    )

    if rels_member in members:
        relationships = ET.fromstring(members[rels_member])
    else:
        relationships = ET.Element(f"{{{rel_ns}}}Relationships")
    relationship_tag = f"{{{rel_ns}}}Relationship"
    for node in list(relationships.findall(relationship_tag)):
        if node.get("Id") == relationship_id:
            relationships.remove(node)
    ET.SubElement(
        relationships,
        relationship_tag,
        {
            "Id": relationship_id,
            "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/font",
            "Target": "fonts/SourceHanSerifCN-Regular.odttf",
        },
    )
    ET.register_namespace("", rel_ns)
    members[rels_member] = ET.tostring(
        relationships, encoding="utf-8", xml_declaration=True
    )
    members[font_member] = font_payload

    settings = ET.fromstring(members["word/settings.xml"])
    if settings.find(f"{{{w_ns}}}embedTrueTypeFonts") is None:
        settings.insert(0, ET.Element(f"{{{w_ns}}}embedTrueTypeFonts"))
    ET.register_namespace("w", w_ns)
    members["word/settings.xml"] = ET.tostring(
        settings, encoding="utf-8", xml_declaration=True
    )

    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(
        prefix=output.stem + "-", suffix=".docx", dir=output.parent, delete=False
    ) as temporary:
        temporary_name = temporary.name
    Path(temporary_name).unlink()
    try:
        with ZipFile(temporary_name, "w", compression=ZIP_DEFLATED) as destination:
            for name, payload in members.items():
                destination.writestr(name, payload)
        shutil.move(temporary_name, output)
    finally:
        try:
            Path(temporary_name).unlink()
        except FileNotFoundError:
            pass


def configure_styles(document: Document) -> None:
    settings = {
        "CUMCM Title": (16, True, WD_ALIGN_PARAGRAPH.CENTER, 0, 12, 0, "Source Han Serif CN"),
        "CUMCM Heading 1": (15, True, WD_ALIGN_PARAGRAPH.CENTER, 12, 8, 0, "Source Han Serif CN"),
        "CUMCM Heading 2": (13, True, WD_ALIGN_PARAGRAPH.LEFT, 9, 5, 0, "Source Han Serif CN"),
        "CUMCM Heading 3": (12, True, WD_ALIGN_PARAGRAPH.LEFT, 7, 4, 0, "Source Han Serif CN"),
        "CUMCM Body": (12, False, WD_ALIGN_PARAGRAPH.JUSTIFY, 0, 0, 24, "Source Han Serif CN"),
        "CUMCM Caption": (10.5, True, WD_ALIGN_PARAGRAPH.CENTER, 4, 6, 0, "Source Han Serif CN"),
    }
    for name, (size, bold, alignment, before, after, first_indent, east_asia) in settings.items():
        style = paragraph_style(document, name)
        set_east_asia_font(style, east_asia)
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.alignment = alignment
        style.paragraph_format.left_indent = Pt(0)
        style.paragraph_format.right_indent = Pt(0)
        style.paragraph_format.first_line_indent = Pt(first_indent)
        style.paragraph_format.line_spacing = 1.25
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = name not in {"CUMCM Body", "CUMCM Caption"}


def add_plain(document: Document, text: str, *, bold: bool = False, centered: bool = False):
    paragraph = document.add_paragraph(style="CUMCM Body")
    paragraph.paragraph_format.first_line_indent = Pt(0) if centered else Pt(24)
    if centered:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        # Placeholder lines are intentionally short; left alignment avoids
        # office renderers stretching them across the full text width.
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text)
    set_run_fonts(run)
    run.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    return paragraph


def build(reference: Path, output: Path, font_zip: Path) -> None:
    if not reference.is_file():
        raise FileNotFoundError(reference)
    document = Document()
    configure_styles(document)

    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.footer_distance = Cm(1.5)
    section.header_distance = Cm(1.5)
    pg_num_type = OxmlElement("w:pgNumType")
    pg_num_type.set(qn("w:start"), "1")
    section._sectPr.append(pg_num_type)
    footer_paragraph = section.footer.paragraphs[0]
    footer_paragraph.clear()
    set_page_number(footer_paragraph)

    document.add_paragraph("论文题目", style="CUMCM Title")
    document.add_paragraph("摘 要", style="CUMCM Heading 1")
    add_plain(document, "摘要正文")
    add_plain(document, "关键词：关键词一；关键词二；关键词三；关键词四", bold=True)
    document.add_page_break()

    structure = (
        ("1 问题重述", ("1.1 问题背景", "1.2 问题回顾")),
        ("2 问题分析", ("2.1 问题一分析",)),
        ("3 模型假设", ()),
        ("4 符号说明", ()),
        ("5 模型建立与求解", ("5.1 问题一的模型建立与求解",)),
        ("6 模型检验与结果分析", ()),
        ("7 模型评价 改进与推广", ("7.1 模型优点", "7.2 模型不足", "7.3 改进与推广")),
    )
    for heading, children in structure:
        document.add_paragraph(heading, style="CUMCM Heading 1")
        for child in children:
            document.add_paragraph(child, style="CUMCM Heading 2")
            add_plain(document, "正文内容")

    document.add_paragraph("AI工具使用声明", style="CUMCM Heading 1")
    add_plain(document, "根据实际使用情况填写声明")
    document.add_paragraph("参考文献", style="CUMCM Heading 1")
    add_plain(document, "按正文首次引用顺序列出真实来源")
    document.add_page_break()
    document.add_paragraph("附录", style="CUMCM Heading 1")
    document.add_paragraph("附录A 支撑材料文件列表", style="CUMCM Heading 2")
    add_plain(document, "列出支撑材料文件及其与论文内容的对应关系")
    document.add_paragraph("附录B 完整源程序", style="CUMCM Heading 2")
    add_plain(document, "放入全部完整可运行代码或明确说明未使用程序")

    # Apply explicit embedded fonts to every run. This prevents office
    # renderers from silently substituting a missing theme font.
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            set_run_fonts(run)

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    embed_source_han(output, font_zip)


def main() -> int:
    parser = argparse.ArgumentParser(description="Build a clean CUMCM Word template")
    parser.add_argument("reference", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument(
        "--font-zip",
        type=Path,
        help="LaTeX template zip containing SourceHanSerifCN-Regular.otf",
    )
    args = parser.parse_args()
    reference = args.reference.expanduser().resolve()
    font_zip = (
        args.font_zip.expanduser().resolve()
        if args.font_zip
        else reference.parent / "2026latex模板.zip"
    )
    build(reference, args.output.expanduser().resolve(), font_zip)
    print(args.output.expanduser().resolve())
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
